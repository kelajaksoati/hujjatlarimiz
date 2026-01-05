from docx import Document
from docx.shared import Pt
import os

def add_watermark(file_path, watermark_text):
    """Word fayliga avtomatik muhr va footer qo'shish"""
    if file_path.endswith('.docx'):
        doc = Document(file_path)
        for section in doc.sections:
            footer = section.footer
            p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            run = p.add_run(f"Manba: {watermark_text} | @ish_reja_uz")
            run.font.size = Pt(9)
        doc.save(file_path)
        return True
    return False

def rename_file(file_path, username):
    """Fayl nomiga kanal userini qo'shish"""
    directory = os.path.dirname(file_path)
    old_name = os.path.basename(file_path)
    new_name = f"@{username}_{old_name}"
    new_path = os.path.join(directory, new_name)
    os.rename(file_path, new_path)
    return new_path, new_name
